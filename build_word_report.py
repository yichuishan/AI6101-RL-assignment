"""Generate an editable Word report from the same content as the two-page PDF."""
from pathlib import Path
import json
from html.parser import HTMLParser
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parent
REPORT_FONT='Times New Roman'
content=json.loads((ROOT/'results/report_content.json').read_text())
doc=Document();section=doc.sections[0]
section.page_width,section.page_height=Mm(210),Mm(297)
section.top_margin=Pt(32);section.bottom_margin=Pt(34)
section.left_margin=section.right_margin=Pt(38)
section.footer_distance=Pt(14)
def set_report_font(font):
    """Set all Word script slots explicitly, removing theme-font overrides."""
    font.name=REPORT_FONT
    rfonts=font._element.get_or_add_rPr().get_or_add_rFonts()
    for slot in ['ascii','hAnsi','eastAsia','cs']:
        rfonts.set(qn('w:'+slot),REPORT_FONT)
    for key in list(rfonts.attrib):
        if key.endswith('Theme'):
            del rfonts.attrib[key]

for style in doc.styles:
    if hasattr(style,'font'):
        set_report_font(style.font)
# Document defaults are outside the named style collection.
for fonts in doc.styles.element.iter(qn('w:rFonts')):
    fonts.attrib.clear()
    for slot in ['ascii','hAnsi','eastAsia','cs']:
        fonts.set(qn('w:'+slot),REPORT_FONT)
normal=doc.styles['Normal'];normal.font.size=Pt(8.3)
normal.font.color.rgb=RGBColor.from_string('233341')
normal.paragraph_format.line_spacing=Pt(11.1)
normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.widow_control=True

class RichText(HTMLParser):
    def __init__(self,paragraph):
        super().__init__(convert_charrefs=True);self.paragraph=paragraph;self.stack=[]
    def handle_starttag(self,tag,attrs):self.stack.append(tag)
    def handle_endtag(self,tag):
        if tag in self.stack:self.stack.remove(tag)
    def handle_data(self,text):
        r=self.paragraph.add_run(text)
        r.bold='b' in self.stack;r.italic='i' in self.stack
        r.font.subscript='sub' in self.stack;r.font.superscript='super' in self.stack

new_page=False
for event in content['events']:
    kind=event['type']
    if kind=='header':
        p=doc.add_paragraph();p.paragraph_format.page_break_before=new_page
        p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=Pt(18)
        r=p.add_run(event['title']);r.bold=True;r.font.size=Pt(14.5)
        r.font.color.rgb=RGBColor.from_string('13334C')
        p=doc.add_paragraph('AI6101 · YourName · YourMatriculationNumber')
        p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(9)
        p.paragraph_format.line_spacing=Pt(10)
        for r in p.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('5A6976')
        borders=OxmlElement('w:pBdr');bottom=OxmlElement('w:bottom')
        for k,v in [('val','single'),('sz','5'),('space','4'),('color','B8CAD8')]:bottom.set(qn('w:'+k),v)
        borders.append(bottom);p._p.get_or_add_pPr().append(borders)
        new_page=False
    elif kind=='block':
        p=doc.add_paragraph();RichText(p).feed(event['text'])
        style=content['styles'][event['kind']]
        p.paragraph_format.line_spacing=Pt(style['leading'])
        p.paragraph_format.space_after=Pt(event['gap'])
        for r in p.runs:r.font.size=Pt(style['font_size'])
        if event['kind']=='head':
            p.paragraph_format.keep_with_next=True
            for r in p.runs:r.bold=True;r.font.color.rgb=RGBColor.from_string('1A4F75')
        elif event['kind']=='small':
            for r in p.runs:r.font.color.rgb=RGBColor.from_string('41515D')
    elif kind=='image':
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(event['gap']);p.paragraph_format.line_spacing=1
        p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(ROOT/'results'/event['name']),width=Pt(event['width']))
    elif kind=='table':
        table=doc.add_table(rows=0,cols=len(event['widths']))
        table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
        for col,width in zip(table.columns,event['widths']):col.width=Pt(width)
        for ri,data in enumerate(event['rows']):
            cells=table.add_row().cells
            for cell,width,text in zip(cells,event['widths'],data):
                cell.width=Pt(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2)
                p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=Pt(9.7)
                RichText(p).feed(str(text))
                for r in p.runs:r.font.size=Pt(7.3)
                if ri==0:
                    for r in p.runs:r.bold=True
                    shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'EAF0F5')
                    cell._tc.get_or_add_tcPr().append(shading)
            # Keep each table row together while retaining editable native cells.
            trpr=table.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        p=doc.add_paragraph();p.paragraph_format.line_spacing=Pt(1);p.paragraph_format.space_after=Pt(3)
        p.add_run().font.size=Pt(1)
    elif kind=='pagebreak':new_page=True

footer=section.footer.paragraphs[0]
footer.paragraph_format.space_after=Pt(0);footer.paragraph_format.line_spacing=Pt(8)
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('AI6101 · ');r.font.size=Pt(7)
for text in ['PAGE',' / ','NUMPAGES']:
    r=footer.add_run();r.font.size=Pt(7)
    if text==' / ':r.add_text(text)
    else:
        begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
        instr=OxmlElement('w:instrText');instr.text=text
        end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
        r._r.extend([begin,instr,end])
doc.core_properties.title='CliffBoxPushing: Q-learning, Exploration and Sparse Rewards'
doc.core_properties.author='YourName — YourMatriculationNumber'
# Explicit run fonts include the title, captions, editable tables and page fields.
for part in [doc.part,*[section.header.part for section in doc.sections],
             *[section.footer.part for section in doc.sections]]:
    for run in part.element.iter(qn('w:r')):
        rpr=run.find(qn('w:rPr'))
        if rpr is None:
            rpr=OxmlElement('w:rPr');run.insert(0,rpr)
        fonts=rpr.find(qn('w:rFonts'))
        if fonts is None:
            fonts=OxmlElement('w:rFonts');rpr.insert(0,fonts)
        fonts.attrib.clear()
        for slot in ['ascii','hAnsi','eastAsia','cs']:
            fonts.set(qn('w:'+slot),REPORT_FONT)
path=ROOT/'deliverables/AI6101_Report_YourName_YourMatriculationNumber.docx'
doc.save(path);print(path)
